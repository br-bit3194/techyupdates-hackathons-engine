"""Comprehensive Unit & Integration Test Suite for TechyUpdates Hackathons Engine."""

import io
import os
import openpyxl
import pytest
from services.ai_extractor import HackathonRecord, fallback_classify_hackathon
from services.collectors.liveness_verifier import (
    generate_dedup_hash,
    is_scam_or_blacklisted,
    is_registration_deadline_active,
    parse_date_heuristic,
    CLOSURE_MARKERS,
)
from services.excel_builder import build_excel_workbook, CATEGORY_CONFIG, HEADERS
from services.telegram_notifier import generate_telegram_caption
from api.trigger import authenticate_request


def test_dedup_hash():
    """Verify MD5 hash generation for platform and title."""
    h1 = generate_dedup_hash("Devpost", "Google AI Hackathon")
    h2 = generate_dedup_hash("DEVPOST ", " google ai hackathon ")
    h3 = generate_dedup_hash("Unstop", "Google AI Hackathon")

    assert h1 == h2, "Hashes should be identical after case/whitespace normalization"
    assert h1 != h3, "Different platforms should produce distinct hashes"


def test_anti_scam_filter():
    """Verify scam and blacklist heuristics."""
    clean_event = {
        "title": "Global Open Source Hackathon",
        "description": "Build decentralized tools for developers.",
        "prize_pool": "$25,000 in Grants",
    }
    scam_event = {
        "title": "Crypto Lottery Hackathon",
        "description": "Send ETH to receive prize. Private key submission required.",
        "prize_pool": "Guaranteed 100x token",
    }

    assert not is_scam_or_blacklisted(clean_event)
    assert is_scam_or_blacklisted(scam_event)


def test_registration_deadline_validation():
    """Verify registration deadline filtering for past, future, and active text."""
    # 1. Past dates should be filtered out
    past_event_1 = {
        "title": "Old Hackathon 2026",
        "registration_deadline": "2026-09-01",
        "description": "Past event",
    }
    past_event_2 = {
        "title": "Concluded Spring Hack",
        "registration_deadline": "Registration Closed",
        "description": "Submissions ended",
    }
    past_event_3 = {
        "title": "Expired Challenge",
        "registration_deadline": "Event Ended",
        "description": "Ended",
    }

    assert is_registration_deadline_active(past_event_1) is False
    assert is_registration_deadline_active(past_event_2) is False
    assert is_registration_deadline_active(past_event_3) is False

    # 2. Future and active dates should be kept
    future_event_1 = {
        "title": "Future AI Hackathon",
        "registration_deadline": "2026-11-20",
        "description": "Upcoming hackathon",
    }
    future_event_2 = {
        "title": "Rolling Sprint",
        "registration_deadline": "Rolling / Open",
        "description": "Active registration",
    }
    future_event_3 = {
        "title": "Ongoing Bounties",
        "registration_deadline": "3 days left",
        "description": "Open now",
    }

    assert is_registration_deadline_active(future_event_1) is True
    assert is_registration_deadline_active(future_event_2) is True
    assert is_registration_deadline_active(future_event_3) is True


def test_hackathon_record_pydantic():
    """Validate strict schema enforcement of HackathonRecord."""
    rec = HackathonRecord(
        platform="Devpost",
        title="Vertex AI Agent Challenge",
        category="AI & GenAI Hackathons",
        theme="Autonomous Agents",
        mode="Online",
        location="Global (Online)",
        prize_pool="$100,000",
        registration_deadline="15 Oct 2026",
        event_dates="20-22 Oct 2026",
        eligibility="All Developers",
        why_participate="Massive prize pool and direct AI mentorship.",
        apply_url="https://devpost.com/hackathons/example",
    )
    assert rec.category == "AI & GenAI Hackathons"
    assert rec.mode == "Online"


def test_fallback_classify_hackathon():
    """Verify heuristic rule classification for all 4 categories."""
    ai_raw = {
        "platform": "Kaggle",
        "title": "LLM Reasoning & Agent Challenge",
        "description": "Build fine-tuned LLM agents.",
        "location": "Global (Online)",
        "prize_pool": "$50,000",
        "apply_url": "https://kaggle.com/c/example",
    }
    web3_raw = {
        "platform": "DoraHacks",
        "title": "Solana High Speed DeFi Hack",
        "description": "Smart contracts and crypto payments.",
        "location": "Online",
        "prize_pool": "$100,000",
        "apply_url": "https://dorahacks.io/example",
    }
    student_raw = {
        "platform": "Major League Hacking (MLH)",
        "title": "HackMIT 2026 Collegiate Hackathon",
        "description": "University hackathon for student builders.",
        "location": "Cambridge, MA",
        "prize_pool": "Swag & Bounties",
        "apply_url": "https://mlh.io/example",
    }
    open_raw = {
        "platform": "Unstop",
        "title": "Flipkart GRiD National Engineering Challenge",
        "description": "Engineering hiring sprint and PPI interviews.",
        "location": "India",
        "prize_pool": "₹15,00,000",
        "apply_url": "https://unstop.com/example",
    }

    ai_rec = fallback_classify_hackathon(ai_raw)
    web3_rec = fallback_classify_hackathon(web3_raw)
    student_rec = fallback_classify_hackathon(student_raw)
    open_rec = fallback_classify_hackathon(open_raw)

    assert ai_rec.category == "AI & GenAI Hackathons"
    assert web3_rec.category == "Web3 & Open Source Hackathons"
    assert student_rec.category == "Student & University Hackathons"
    assert open_rec.category == "Open Innovation & Hiring Challenges"


def test_excel_workbook_structure():
    """Validate 4-tab workbook generation, formatting, headers, and clickable links."""
    sample_records = [
        HackathonRecord(
            platform="Devpost",
            title="Google Gemini AI Sprint",
            category="AI & GenAI Hackathons",
            theme="GenAI Agents",
            mode="Online",
            location="Global (Online)",
            prize_pool="$100,000",
            registration_deadline="15 Oct 2026",
            event_dates="20-22 Oct 2026",
            eligibility="All Developers",
            why_participate="Direct API credits and cash prizes.",
            apply_url="https://devpost.com/hackathons/example",
        ),
        HackathonRecord(
            platform="Devfolio",
            title="ETHIndia 2026",
            category="Web3 & Open Source Hackathons",
            theme="Ethereum Layer-2",
            mode="In-Person",
            location="Bengaluru, India",
            prize_pool="$150,000",
            registration_deadline="01 Nov 2026",
            event_dates="04-06 Dec 2026",
            eligibility="Web3 Developers",
            why_participate="Asia's biggest Ethereum hackathon.",
            apply_url="https://ethindia.devfolio.co",
        ),
    ]

    buffer = build_excel_workbook(sample_records)
    assert isinstance(buffer, io.BytesIO)
    assert buffer.getbuffer().nbytes > 0

    wb = openpyxl.load_workbook(buffer)
    expected_sheets = [cfg["sheet_name"] for cfg in CATEGORY_CONFIG]
    assert wb.sheetnames == expected_sheets

    # Verify sheet 1 headers
    ws = wb[expected_sheets[0]]
    assert ws.cell(row=1, column=1).value == "Organizer / Platform"
    assert ws.cell(row=1, column=11).value == "Direct Apply Link"
    assert ws.freeze_panes == "A2"

    # Verify link formula
    link_formula = ws.cell(row=2, column=11).value
    assert link_formula.startswith('=HYPERLINK("https://devpost.com/hackathons/example"')


def test_telegram_caption_generation():
    """Verify humanized TechyUpdates community caption copy."""
    records = [
        HackathonRecord(
            platform="Google",
            title="Global Gemini AI Challenge",
            category="AI & GenAI Hackathons",
            theme="GenAI",
            mode="Online",
            location="Global (Online)",
            prize_pool="$100,000",
            registration_deadline="15 Oct 2026",
            event_dates="20 Oct 2026",
            eligibility="All Developers",
            why_participate="Huge prizes",
            apply_url="https://example.com",
        )
    ]
    caption = generate_telegram_caption(records)
    assert "TechyUpdates" in caption
    assert "*AI & GenAI Hackathons:* 1" in caption
    assert "Attached Excel file" in caption


def test_authenticate_request():
    """Validate Bearer and query parameter authentication."""
    os.environ["CRON_SECRET"] = "secret_12345"

    assert authenticate_request({"Authorization": "Bearer secret_12345"}) is True
    assert authenticate_request({"Authorization": "Bearer wrong_secret"}) is False
    assert authenticate_request({}, {"secret": ["secret_12345"]}) is True
    assert authenticate_request({"User-Agent": "vercel-cron/1.0"}) is True
    assert authenticate_request({}) is False


def test_liveness_closure_markers():
    """Validate detection of closed registration markers."""
    html_closed = "<html><body>Sorry, registration closed yesterday.</body></html>"
    html_open = "<html><body>Register now for the hackathon! Form is live.</body></html>"

    has_closed_marker = any(m in html_closed.lower() for m in CLOSURE_MARKERS)
    assert has_closed_marker is True

    has_closed_in_open = any(m in html_open.lower() for m in CLOSURE_MARKERS)
    assert has_closed_in_open is False


@pytest.mark.asyncio
async def test_hack2skill_collector():
    """Verify Hack2skill collector integration."""
    from services.collectors.hack2skill import fetch_hack2skill_hackathons
    import httpx
    async with httpx.AsyncClient() as client:
        results = await fetch_hack2skill_hackathons(client)
        assert isinstance(results, list)
        assert len(results) > 0
        first = results[0]
        assert "Hack2skill" in first["platform"]
        assert first["title"]
        assert first["apply_url"].startswith("http")

